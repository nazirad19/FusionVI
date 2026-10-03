"""Build the concise FusionVI paper-benchmark technical report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "FusionVI_Paper_Benchmark_Technical_Report.docx"
NAVY = "172554"
ORANGE = "F97316"
LIGHT_BLUE = "EFF6FF"
BORDER = "D9D9D9"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths: list[float] | None = None) -> None:
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if widths:
                cell.width = Inches(widths[c_idx])
            if r_idx == 0:
                set_cell_fill(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True
            elif r_idx % 2 == 0:
                set_cell_fill(cell, LIGHT_BLUE)
            else:
                set_cell_fill(cell, "FFFFFF")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = 1
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def metric_sentence(delta: float, total: float, fusion: float, parameter_reduction: float) -> str:
    relative = 100.0 * (total - fusion) / total
    if delta < -0.005:
        return (
            f"FusionVI reduced mean RMSLE from {total:.4f} to {fusion:.4f} "
            f"({relative:.2f}% lower) while using {parameter_reduction:.1f}% fewer trainable parameters."
        )
    if abs(delta) <= 0.005:
        direction = "lower" if delta < 0 else "higher"
        return (
            f"The models were nearly tied: FusionVI RMSLE was {fusion:.4f} versus {total:.4f} for totalVI "
            f"({abs(delta):.4f} {direction}). FusionVI used {parameter_reduction:.1f}% fewer trainable parameters."
        )
    return (
        f"FusionVI RMSLE was {fusion:.4f} versus {total:.4f} for totalVI "
        f"({delta:.4f} higher) despite using {parameter_reduction:.1f}% fewer trainable parameters."
    )


def main() -> None:
    headline = json.loads((ROOT / "results" / "paper_benchmark_headline.json").read_text())
    seeds = pd.read_csv(ROOT / "results" / "paper_benchmark_seed_summary.csv")
    data = json.loads((ROOT / "results" / "paper_benchmark_data_summary.json").read_text())
    complete = headline["completed_runs"]
    params = {run["model"]: run["trainable_parameters"] for run in complete}
    total_params = params["totalVI"]
    fusion_params = params["FusionVI"]
    parameter_reduction = 100.0 * (total_params - fusion_params) / total_params
    total = headline["totalvi_mean_rmsle"]
    fusion = headline["fusionvi_mean_rmsle"]
    delta = headline["fusionvi_minus_totalvi_rmsle"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 29), ("Heading 1", 17), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)
    styles["List Bullet"].font.name = "Arial"
    styles["List Bullet"].font.size = Pt(10.5)
    styles["List Bullet"].paragraph_format.space_after = Pt(2)

    # Header and footer.
    header = sec.header.paragraphs[0]
    header.text = "FusionVI  |  Paper-aligned benchmark"
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor(71, 85, 105)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Multimodal AI for therapeutic discovery")
    footer.runs[0].font.name = "Arial"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    title = doc.add_paragraph()
    title.style = styles["Title"]
    title.add_run("FusionVI")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    run = subtitle.add_run("A paper-aligned benchmark for missing-protein recovery in CITE-seq")
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(23, 37, 84)
    meta = doc.add_paragraph()
    meta.add_run("Technical report  •  totalVI versus FusionVI  •  October 2026").bold = True
    meta.paragraph_format.space_after = Pt(18)

    doc.add_heading("Abstract", level=1)
    abstract = (
        "Therapeutic biomarker studies often combine sparse transcript counts with surface-protein measurements, "
        "but entire protein panels may be unavailable in a new batch. We reproduced the totalVI paper's Figure 3 "
        "missing-protein design on the official SLN111 CITE-seq object and compared the published totalVI architecture "
        "with FusionVI, which replaces the joint encoder with separate RNA and protein branches and a learned cell-level gate. "
        f"Across {headline['completed_totalvi_initializations']} paired random initializations and 110 proteins, "
        + metric_sentence(delta, total, fusion, parameter_reduction)
        + " This course-scale benchmark estimates algorithmic performance; it does not establish population-level biological generalization."
    )
    doc.add_paragraph(abstract)

    doc.add_heading("1. Biological and technical question", level=1)
    doc.add_paragraph(
        "Can RNA measured in a new biological batch, together with multimodal structure learned from a reference batch, "
        "recover an entirely unmeasured surface-protein panel? This matters for therapeutic discovery because surface proteins "
        "define immune populations, pharmacodynamic states and drug targets, while antibody panels can fail or differ between experiments."
    )
    doc.add_paragraph(
        "The model question was narrower: does modality-specific encoding improve the original totalVI decoder under the exact "
        "missing-panel stress test used in the paper?"
    )

    doc.add_heading("2. Dataset and study design", level=1)
    add_table(
        doc,
        ["Component", "Value"],
        [
            ["Source", "Official SLN111 CITE-seq object released with totalVI (GSE150599)"],
            ["Cells", f"{data['cells']:,}"],
            ["RNA features", f"{data['genes']:,} highly variable genes"],
            ["Protein features", f"{data['proteins']} biological proteins after HTO removal"],
            ["Reference batch", f"SLN111-D1, {data['source_cells']:,} cells, RNA and proteins retained"],
            ["Target batch", f"SLN111-D2, {data['target_cells']:,} cells, all protein inputs hidden"],
        ],
        [1.55, 5.25],
    )
    doc.add_paragraph(
        "Untouched D2 protein counts were stored only as evaluation truth. During training, totalVI's missing-protein mask removed "
        "the entire D2 panel from the protein likelihood. Both models saw the same cells, gene features, batch labels and validation split."
    )

    doc.add_heading("3. Models", level=1)
    doc.add_heading("Published totalVI baseline", level=2)
    doc.add_paragraph(
        "The baseline uses the scvi-tools totalVI implementation: a joint RNA-protein encoder, a 20-dimensional latent state, "
        "a negative-binomial RNA likelihood and the native background/foreground protein mixture decoder."
    )
    doc.add_heading("FusionVI", level=2)
    doc.add_paragraph(
        "FusionVI keeps totalVI's generative decoder and likelihoods. It changes only the inference network: one branch encodes RNA, "
        "a second branch encodes protein measurements, and a learned gate combines their hidden states for each cell before estimating "
        "the same 20-dimensional latent distribution. This separates modality-specific representations while preserving the paper's output model."
    )
    add_table(
        doc,
        ["Model", "Encoder", "Decoder and likelihood", "Trainable parameters"],
        [
            ["totalVI", "Joint RNA-protein encoder", "Native totalVI", f"{total_params:,}"],
            ["FusionVI", "RNA branch + protein branch + cell gate", "Native totalVI", f"{fusion_params:,}"],
        ],
        [1.05, 2.35, 1.75, 1.65],
    )

    doc.add_heading("4. Training and evaluation", level=1)
    add_bullet(doc, "Five paired seeds: 2026 to 2030. Each seed initializes both models independently.")
    add_bullet(doc, "Maximum 500 epochs, Adam learning rate 0.004, batch size 256, 90/10 train-validation split.")
    add_bullet(doc, "Early stopping patience 45 with learning-rate reduction on validation plateau.")
    add_bullet(doc, "Predictions average 25 posterior samples and decode D2 cells through the D1 batch.")
    add_bullet(doc, "Primary metric: mean per-protein RMSLE across 7,564 D2 cells. Lower values indicate better reconstruction.")
    doc.add_paragraph(
        "The paper used 30 initializations. We used five paired initializations to fit the course compute budget, so this is a "
        "protocol-aligned reproduction rather than an exact reproduction of the paper's uncertainty analysis."
    )

    doc.add_heading("5. Results", level=1)
    doc.add_paragraph(metric_sentence(delta, total, fusion, parameter_reduction))
    seed_pivot = seeds.pivot(index="seed", columns="model", values="mean_rmsle").reset_index()
    seed_rows = []
    for _, row in seed_pivot.iterrows():
        diff = row["FusionVI"] - row["totalVI"]
        seed_rows.append([str(int(row["seed"])), f"{row['totalVI']:.4f}", f"{row['FusionVI']:.4f}", f"{diff:+.4f}"])
    seed_rows.append(["Mean", f"{total:.4f}", f"{fusion:.4f}", f"{delta:+.4f}"])
    add_table(doc, ["Seed", "totalVI RMSLE", "FusionVI RMSLE", "FusionVI − totalVI"], seed_rows, [1.1, 1.75, 1.75, 2.15])

    figure = ROOT / "results" / "figures" / "paper_benchmark_totalvi_vs_fusionvi.png"
    if figure.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(figure), width=Inches(6.75))
        cap = doc.add_paragraph(
            "Figure 1. Paired seed means, per-protein errors and the distribution of paired differences. Negative differences favor FusionVI."
        )
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(8.5)

    wins = headline["proteins_fusionvi_better"]
    n_proteins = headline["proteins_compared"]
    doc.add_paragraph(
        f"FusionVI had lower mean RMSLE for {wins} of {n_proteins} proteins. A paired Wilcoxon test across protein-level mean errors "
        f"gave p = {headline['protein_level_paired_wilcoxon_p']:.3g}. Proteins and seeds are algorithmic benchmark units, so this test "
        "describes the benchmark and must not be interpreted as evidence from independent animals or patients."
    )

    doc.add_heading("Secondary four-marker extension", level=2)
    doc.add_paragraph(
        "A separate leave-one-mouse-out extension hid CD20, CD28, CD4 and CD8a while retaining the other 106 proteins. "
        "FusionVI achieved a mean Spearman correlation of 0.673 versus 0.578 for totalVI. This result tests marker ranking in two held-out "
        "mice and uses a different readout, so it complements rather than replaces the paper-aligned benchmark."
    )

    doc.add_heading("6. Interpretation", level=1)
    if delta < 0:
        doc.add_paragraph(
            "The full-panel result supports the claim that modality-specific encoding can match or modestly improve totalVI's missing-protein "
            "reconstruction while reducing the parameter count. The result does not prove that FusionVI is universally better: the effect is "
            "small, the benchmark contains only one source-target pair and the paper's 30-initialization analysis was not repeated in full."
        )
    else:
        doc.add_paragraph(
            "The full-panel result does not support a superiority claim for FusionVI on the paper's primary missing-protein metric. "
            "Its smaller encoder may still offer parameter efficiency, and the positive four-marker result suggests value for targeted panel recovery."
        )
    doc.add_paragraph(
        "Biologically, successful reconstruction means that RNA and the reference donor encode enough shared immune-cell structure to estimate "
        "many unmeasured surface markers. The experiment evaluates technical transfer and panel completion. It does not establish therapeutic "
        "response prediction, clinical utility or generalization across a population."
    )

    doc.add_heading("7. Reproducibility", level=1)
    add_bullet(doc, "Primary scripts: prepare_paper_benchmark.py, train_paper_benchmark.py and evaluate_paper_benchmark.py.")
    add_bullet(doc, "Configuration: config/paper_benchmark.yaml records data dimensions, training settings, seeds and the paper's 30-run reference.")
    add_bullet(doc, "Saved outputs include every per-protein metric, seed summary, paired differences, figures and completion metadata.")
    add_bullet(doc, "Raw data and model weights are regenerated locally and excluded from Git; source code and compact results are versioned.")
    doc.add_paragraph("Run from PowerShell:")
    code = doc.add_paragraph()
    code.style = styles["Normal"]
    code.paragraph_format.left_indent = Inches(0.3)
    r = code.add_run(".\\run_paper_benchmark.ps1")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    doc.add_heading("References", level=1)
    doc.add_paragraph(
        "Gayoso A, Steier Z, Lopez R, et al. Joint probabilistic modeling of single-cell multi-omic data with totalVI. "
        "Nature Methods. 2021;18:272–282. https://doi.org/10.1038/s41592-020-01050-x"
    )
    doc.add_paragraph(
        "Yosef Lab. totalVI reproducibility repository. https://github.com/YosefLab/totalVI_reproducibility"
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
